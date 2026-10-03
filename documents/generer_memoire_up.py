# -*- coding: utf-8 -*-
"""Génère le mémoire PLANIPRO en Word (.docx) — méthodologie Processus Unifié (UP)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent / "MEMOIRE_PLANIPRO.docx"

doc = Document()
for sec in doc.sections:
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)


def font(run, size=12, bold=False, italic=False, name="Times New Roman", color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def p(text="", *, center=False, justify=True, bold=False, italic=False, size=12,
      first_indent=None, space_before=0, space_after=8, line=1.5, indent_left=0):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justify:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if indent_left:
        pf.left_indent = Cm(indent_left)
    if text:
        r = para.add_run(text)
        font(r, size=size, bold=bold, italic=italic)
    return para


def rich(parts, *, center=False, justify=True, first_indent=1.25, size=12, space_after=8):
    """parts: list of (text, bold, italic) or str"""
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justify:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_indent:
        pf.first_line_indent = Cm(first_indent)
    for part in parts:
        if isinstance(part, str):
            r = para.add_run(part)
            font(r, size=size)
        else:
            t, b, i = part[0], part[1] if len(part) > 1 else False, part[2] if len(part) > 2 else False
            r = para.add_run(t)
            font(r, size=size, bold=b, italic=i)
    return para


def h1(text):
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(WD_BREAK.PAGE)
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(14)
    r = para.add_run(text.upper())
    font(r, size=16, bold=True)


def h2(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(14)
    para.paragraph_format.space_after = Pt(8)
    r = para.add_run(text)
    font(r, size=14, bold=True)


def h3(text):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(6)
    r = para.add_run(text)
    font(r, size=12, bold=True, italic=True)


def bullet(text, ordered=False, n=1):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(1.2)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.3
    prefix = f"{n}. " if ordered else "• "
    r = para.add_run(prefix + text)
    font(r, size=12)


def figure(num, caption, hint):
    p(f"[ EMPLACEMENT FIGURE {num} ]", center=True, justify=False, bold=True, size=11,
      first_indent=0, space_before=12, space_after=4)
    p(hint, center=True, justify=False, italic=True, size=10, first_indent=0, space_after=4)
    p(f"Figure {num} — {caption}", center=True, justify=False, italic=True, size=11,
      first_indent=0, space_after=12)
    p("─" * 42, center=True, justify=False, size=10, first_indent=0, space_after=8)


def table(headers, rows, caption=None):
    if caption:
        p(caption, center=True, justify=False, italic=True, size=11, first_indent=0, space_after=6)
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        rp = cell.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = rp.add_run(h)
        font(run, size=10, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            font(run, size=10)
    doc.add_paragraph()


def ref(text):
    para = doc.add_paragraph()
    pf = para.paragraph_format
    pf.left_indent = Cm(1.0)
    pf.first_line_indent = Cm(-1.0)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.15
    r = para.add_run(text)
    font(r, size=11)


def page():
    doc.add_page_break()


# =============================================================================
# PAGE DE GARDE
# =============================================================================
p("[NOM DE L'UNIVERSITÉ / INSTITUT]", center=True, justify=False, bold=True, size=14, first_indent=0)
p("[Faculté / Département d'Informatique]", center=True, justify=False, size=12, first_indent=0)
p("République Démocratique du Congo", center=True, justify=False, size=11, first_indent=0)
p("")
p("[LOGO INSTITUTION]", center=True, justify=False, italic=True, size=11, first_indent=0)
p("")
p("MÉMOIRE DE FIN D'ÉTUDES", center=True, justify=False, bold=True, size=12, first_indent=0)
p("")
p("Conception et réalisation d'une application web de gestion des rendez-vous et des activités professionnelles : cas du système PLANIPRO",
  center=True, justify=False, bold=True, size=14, first_indent=0)
p("")
p("Présenté en vue de l'obtention du grade de [Licence / Graduat] en Informatique",
  center=True, justify=False, italic=True, size=11, first_indent=0)
p("Option : [Génie Logiciel / Systèmes d'information]", center=True, justify=False, size=11, first_indent=0)
p("Méthodologie : Processus Unifié (Unified Process — UP)", center=True, justify=False, bold=True, size=11, first_indent=0)
p("")
p("Présenté et défendu par :", center=True, justify=False, bold=True, size=12, first_indent=0)
p("[NOM COMPLET DE L'ÉTUDIANT(E)]", center=True, justify=False, size=12, first_indent=0)
p("")
p("Sous la direction de :", center=True, justify=False, bold=True, size=12, first_indent=0)
p("[NOM DU PROMOTEUR]", center=True, justify=False, size=12, first_indent=0)
p("[Grade académique]", center=True, justify=False, size=10, first_indent=0)
p("")
p("Année académique : 2024 – 2025", center=True, justify=False, bold=True, size=12, first_indent=0)

# Épigraphe
page()
p("« La meilleure façon de prédire l'avenir est de le concevoir. »", center=True, justify=False, italic=True, size=12, first_indent=0)
p("— Alan Kay", center=True, justify=False, italic=True, size=11, first_indent=0)

# Dédicace
page()
p("DÉDICACE", center=True, justify=False, bold=True, size=16, first_indent=0)
p("")
p("À Dieu Tout-Puissant, source de toute sagesse.", first_indent=1.25)
p("À mes parents, pour leur amour et leurs sacrifices.", first_indent=1.25)
p("À mes frères, sœurs et proches, pour leur soutien permanent.", first_indent=1.25)
p("À tous ceux qui ont contribué à la réalisation de ce travail.", first_indent=1.25)
p("")
p("[Prénom de l'auteur]", center=True, justify=False, bold=True, first_indent=0)

# Remerciements
page()
p("REMERCIEMENTS", center=True, justify=False, bold=True, size=16, first_indent=0)
p("Au terme de ce mémoire, nous exprimons notre profonde gratitude à toutes les personnes qui ont contribué à son aboutissement.", first_indent=1.25)
p("Nous remercions particulièrement notre promoteur, [Nom du promoteur], pour la qualité de son encadrement scientifique, sa disponibilité et la rigueur méthodologique qu'il a su nous transmettre, notamment dans l'application du Processus Unifié.", first_indent=1.25)
p("Nos remerciements s'adressent également au corps professoral de [Institution], dont les enseignements en génie logiciel, bases de données et systèmes d'information ont constitué le socle de ce travail.", first_indent=1.25)
p("Nous remercions les professionnels qui ont partagé leur expérience sur la gestion des rendez-vous et des activités, alimentant ainsi l'étude de l'existant.", first_indent=1.25)
p("Enfin, nous adressons un merci particulier à notre famille et à nos camarades pour leur soutien moral constant.", first_indent=1.25)

# Résumé
page()
p("RÉSUMÉ", center=True, justify=False, bold=True, size=16, first_indent=0)
p("La digitalisation des processus organisationnels constitue un enjeu majeur pour les entreprises. La gestion des rendez-vous professionnels, le suivi des activités internes et la communication avec les clients, partenaires et visiteurs restent souvent fragmentées : agendas papier, tableurs, appels et messageries personnelles non tracées.", first_indent=0)
p("Ce mémoire présente la conception et la réalisation de PLANIPRO, une application web multi-rôles (entreprise, visiteur, client, partenaire) permettant de centraliser les rendez-vous, les activités, les contacts, la messagerie, le calendrier, les notifications et les rapports. Le développement a été conduit selon le Processus Unifié (Unified Process — UP), structuré en phases d'Inception, d'Élaboration, de Construction et de Transition, et organisé autour des disciplines de modélisation métier, exigences, analyse et conception, implémentation et tests.", first_indent=0)
p("La stack technique repose sur PHP 8+, MySQL/MariaDB, Apache (XAMPP), HTML/CSS/JavaScript et FullCalendar. Les résultats montrent un prototype fonctionnel couvrant le cycle de vie des rendez-vous, le suivi des activités, la messagerie intégrée et la sécurité de base (sessions, rôles, CSRF, hachage des mots de passe).", first_indent=0)
rich([("Mots-clés : ", True, False), "Processus Unifié, UP, génie logiciel, système d'information, gestion des rendez-vous, activités professionnelles, application web, PHP, MySQL, cas d'utilisation, UML."], first_indent=0)

p("")
p("ABSTRACT", center=True, justify=False, bold=True, size=14, first_indent=0)
p("Organizational digitalization is a major challenge for enterprises. Appointment management, internal activity tracking and communication with clients, partners and visitors are still often fragmented.", first_indent=0)
p("This thesis presents the design and implementation of PLANIPRO, a multi-role web application that centralizes appointments, activities, contacts, messaging, calendar, notifications and reporting. Development followed the Unified Process (UP) through Inception, Elaboration, Construction and Transition phases.", first_indent=0)
p("The technical stack is PHP 8+, MySQL/MariaDB, Apache (XAMPP), HTML/CSS/JavaScript and FullCalendar. The resulting prototype supports the appointment lifecycle, activity tracking, integrated messaging and baseline security.", first_indent=0)
rich([("Keywords: ", True, False), "Unified Process, UP, software engineering, information system, appointment management, professional activities, web application, PHP, MySQL, use cases, UML."], first_indent=0)

# TOC
page()
p("TABLE DES MATIÈRES", center=True, justify=False, bold=True, size=16, first_indent=0)
toc_items = [
    ("INTRODUCTION GÉNÉRALE", True),
    ("0.1 Contexte et justification", False),
    ("0.2 Problématique", False),
    ("0.3 Hypothèses de recherche", False),
    ("0.4 Objectifs de l'étude", False),
    ("0.5 Intérêt du sujet", False),
    ("0.6 Délimitation du sujet", False),
    ("0.7 Méthodologie : le Processus Unifié (UP)", False),
    ("0.8 Subdivision du travail", False),
    ("CHAPITRE I — CADRE THÉORIQUE ET CONCEPTUEL", True),
    ("1.1 Introduction du chapitre", False),
    ("1.2 Les systèmes d'information", False),
    ("1.3 Gestion des rendez-vous et du temps organisationnel", False),
    ("1.4 Gestion des activités et des tâches", False),
    ("1.5 Applications web et architectures logicielles", False),
    ("1.6 Bases de données relationnelles", False),
    ("1.7 Sécurité des applications web", False),
    ("1.8 Le Processus Unifié (UP) : fondements et principes", False),
    ("1.9 UML et artefacts du Processus Unifié", False),
    ("1.10 Concepts clés de PLANIPRO", False),
    ("1.11 Conclusion du chapitre", False),
    ("CHAPITRE II — ÉTUDE DE L'EXISTANT ET ANALYSE DU PROBLÈME", True),
    ("2.1 Introduction (phase d'Inception)", False),
    ("2.2 Présentation du domaine d'étude", False),
    ("2.3 Modélisation métier de l'existant (as-is)", False),
    ("2.4 Analyse critique et risques", False),
    ("2.5 Formulation du problème et vision du système", False),
    ("2.6 Analyse comparative des solutions", False),
    ("2.7 Besoins et exigences initiales", False),
    ("2.8 Conclusion du chapitre", False),
    ("CHAPITRE III — ANALYSE ET CONCEPTION DU SYSTÈME", True),
    ("3.1 Introduction (phase d'Élaboration)", False),
    ("3.2 Exigences fonctionnelles et non fonctionnelles", False),
    ("3.3 Acteurs et cas d'utilisation (UP)", False),
    ("3.4 Scénarios et spécifications des CU", False),
    ("3.5 Analyse (modèle d'analyse)", False),
    ("3.6 Conception de la base de données", False),
    ("3.7 Architecture logicielle (vues 4+1)", False),
    ("3.8 Conception des interfaces", False),
    ("3.9 Sécurité et règles de gestion", False),
    ("3.10 Conclusion du chapitre", False),
    ("CHAPITRE IV — IMPLÉMENTATION DU SYSTÈME PROPOSÉ", True),
    ("4.1 Introduction (Construction et Transition)", False),
    ("4.2 Environnement et organisation du code", False),
    ("4.3 Itérations de construction", False),
    ("4.4 Implémentation des modules", False),
    ("4.5 Interfaces réalisées", False),
    ("4.6 Tests et transition", False),
    ("4.7 Difficultés, limites et perspectives", False),
    ("4.8 Conclusion du chapitre", False),
    ("CONCLUSION GÉNÉRALE ET PERSPECTIVES", True),
    ("BIBLIOGRAPHIE ET WEBOGRAPHIE", True),
    ("ANNEXES", True),
]
for item, is_ch in toc_items:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.2
    if not is_ch:
        para.paragraph_format.left_indent = Cm(0.75)
    r = para.add_run(item)
    font(r, size=11, bold=is_ch)

# Listes figures / tableaux / sigles
page()
p("LISTE DES FIGURES", center=True, justify=False, bold=True, size=16, first_indent=0)
figs = [
    "0.1 Logo de l'institution",
    "0.2 Fragmentation des outils vs PLANIPRO",
    "0.3 Cycle de vie UP (phases et disciplines)",
    "1.1 Pyramide des SI",
    "1.2 Architecture web à trois tiers",
    "1.3 Modèle MVC / pages-actions",
    "1.4 Cycle de vie d'un rendez-vous",
    "1.5 Phases du Processus Unifié",
    "1.6 Disciplines UP et artefacts",
    "1.7 Diagramme de cas d'utilisation (principe UP)",
    "2.1 Cartographie des acteurs du domaine",
    "2.2 Processus as-is de prise de RDV",
    "2.3 Diagramme d'Ishikawa",
    "2.4 Matrice SWOT",
    "2.5 Comparaison des outils du marché",
    "2.6 Document de Vision (structure UP)",
    "3.1 Modèle de cas d'utilisation global PLANIPRO",
    "3.2 CU — module Rendez-vous",
    "3.3 CU — module Activités",
    "3.4 Diagramme de séquence : validation RDV",
    "3.5 Diagramme d'activités : contact public → visiteur",
    "3.6 Modèle d'analyse (classes d'analyse)",
    "3.7 MCD / diagramme entité-association",
    "3.8 MLD relationnel",
    "3.9 Architecture 4+1 de Kruchten",
    "3.10 Maquette accueil public",
    "3.11 Maquette tableau de bord entreprise",
    "3.12 Maquette calendrier",
    "3.13 Maquette messagerie",
    "4.1 Capture : accueil PLANIPRO",
    "4.2 Capture : inscription / connexion",
    "4.3 Capture : tableau de bord entreprise",
    "4.4 Capture : gestion des rendez-vous",
    "4.5 Capture : activités",
    "4.6 Capture : calendrier",
    "4.7 Capture : messagerie",
    "4.8 Capture : rapports PDF",
    "4.9 Capture : espace client",
    "4.10 Capture : notifications publiques",
    "4.11 Schéma de déploiement XAMPP",
    "4.12 Plan d'itérations UP (Construction)",
]
for i, f in enumerate(figs, 1):
    bullet(f"Figure : {f}")

page()
p("LISTE DES TABLEAUX", center=True, justify=False, bold=True, size=16, first_indent=0)
tabs = [
    "1.1 Types de systèmes d'information",
    "1.2 Comparaison des approches de développement",
    "1.3 Phases et objectifs du Processus Unifié",
    "1.4 Artefacts UP mobilisés dans PLANIPRO",
    "2.1 Inventaire des pratiques actuelles",
    "2.2 Forces et faiblesses de l'existant",
    "2.3 Comparaison fonctionnelle des solutions",
    "2.4 Besoins fonctionnels prioritaires",
    "2.5 Besoins non fonctionnels",
    "2.6 Matrice des risques",
    "3.1 Acteurs du système",
    "3.2 Matrice des droits d'accès",
    "3.3 Dictionnaire de données (extrait RDV)",
    "3.4 Relations et cardinalités",
    "3.5 Règles de gestion",
    "3.6 Exigences non fonctionnelles détaillées",
    "4.1 Environnement logiciel",
    "4.2 Structure des répertoires",
    "4.3 Itérations de Construction UP",
    "4.4 Scénarios de test",
    "4.5 Couverture fonctionnelle",
    "4.6 Passage as-is → to-be",
]
for t in tabs:
    bullet(f"Tableau : {t}")

page()
p("LISTE DES ABRÉVIATIONS ET SIGLES", center=True, justify=False, bold=True, size=16, first_indent=0)
table(
    ["Sigle", "Signification"],
    [
        ["UP", "Unified Process (Processus Unifié)"],
        ["RUP", "Rational Unified Process"],
        ["UML", "Unified Modeling Language"],
        ["CU", "Cas d'utilisation (Use Case)"],
        ["SI", "Système d'Information"],
        ["MVC", "Model-View-Controller"],
        ["MCD/MLD/MPD", "Modèles conceptuel / logique / physique de données"],
        ["PDO", "PHP Data Objects"],
        ["CSRF", "Cross-Site Request Forgery"],
        ["XSS", "Cross-Site Scripting"],
        ["CRUD", "Create, Read, Update, Delete"],
        ["PRG", "Post-Redirect-Get"],
        ["RDV", "Rendez-vous"],
        ["SGBD", "Système de Gestion de Bases de Données"],
        ["XAMPP", "Apache + MariaDB/MySQL + PHP + Perl"],
        ["OWASP", "Open Web Application Security Project"],
        ["PMI", "Project Management Institute"],
    ],
)

# =============================================================================
# INTRODUCTION
# =============================================================================
h1("Introduction générale")

h2("0.1 Contexte et justification")
p("Depuis plusieurs décennies, les Technologies de l'Information et de la Communication (TIC) transforment le fonctionnement des organisations. L'information est devenue un facteur critique de compétitivité (Laudon & Laudon, 2020 ; Reix et al., 2016). Dans ce contexte, la capacité d'une entreprise à planifier, coordonner et suivre ses interactions professionnelles — rendez-vous, activités, échanges avec clients et partenaires — conditionne la qualité de service et la productivité (Mintzberg, 1994 ; Drucker, 2006).")
p("Or, de nombreuses PME continuent de s'appuyer sur des outils disparates : agendas papier, tableurs, appels téléphoniques, messageries instantanées personnelles. Ces pratiques génèrent pertes d'information, doubles réservations, oublis de suivi et absence de traçabilité (Davenport, 1993 ; Hammer & Champy, 2003).")
p("Parallèlement, le génie logiciel dispose de méthodes éprouvées pour conduire le développement de systèmes complexes. Parmi elles, le Processus Unifié (Unified Process — UP), popularisé notamment sous la forme du Rational Unified Process (RUP), propose une démarche itérative, incrémentale, pilotée par les cas d'utilisation et centrée sur l'architecture (Jacobson et al., 1999 ; Kruchten, 2003 ; Larman, 2004). C'est cette méthodologie que nous avons retenue pour concevoir et réaliser PLANIPRO.")
p("PLANIPRO est une application web permettant à une entreprise de recevoir et valider des rendez-vous, d'organiser des activités internes, de gérer clients, partenaires et visiteurs, de communiquer via une messagerie, de visualiser un calendrier, de suivre des notifications et de produire des rapports.")
figure("0.2", "Du chaos informationnel à la plateforme PLANIPRO",
       "À dessiner : outils fragmentés (agenda, Excel, WhatsApp, téléphone) → flèche → plateforme centralisée PLANIPRO.")

h2("0.2 Problématique")
p("Malgré l'existence d'outils grand public (Google Calendar, Calendly, Outlook, etc.), beaucoup d'entreprises peinent à disposer d'une solution adaptée à leur réalité : multi-acteurs, validation formelle des RDV, suivi des activités, messagerie métier, rapports et isolation des données par organisation (Gartner, 2022).")
p("La question centrale est donc :", first_indent=0)
p("Comment concevoir et réaliser, selon le Processus Unifié (UP), une application web permettant à une entreprise de centraliser la gestion de ses rendez-vous, de ses activités professionnelles et de ses échanges avec les clients, partenaires et visiteurs, tout en garantissant la traçabilité, la sécurité et la simplicité d'usage ?",
  italic=True, first_indent=0, indent_left=1.0)
p("Questions secondaires :", first_indent=0)
for i, q in enumerate([
    "Quels fondements théoriques (SI, web, UP, UML) mobiliser ?",
    "Quelles limites de l'existant et quels risques métier ?",
    "Quelles exigences et quels cas d'utilisation prioriser en Élaboration ?",
    "Quelle architecture et quel modèle de données retenir ?",
    "Comment construire, tester et préparer la Transition du système ?",
], 1):
    bullet(q, ordered=True, n=i)

h2("0.3 Hypothèses de recherche")
rich([("Hypothèse principale (H0) : ", True, False),
      "Le développement itératif et incrémental de PLANIPRO selon le Processus Unifié, piloté par les cas d'utilisation, permet d'obtenir un système web multi-rôles qui améliore la coordination des rendez-vous et des activités par rapport aux pratiques fragmentées."], first_indent=0)
p("Hypothèses secondaires :", first_indent=0)
bullet("H1 : La formalisation du cycle de vie des RDV (en attente, accepté, refusé, reporté, annulé, terminé) réduit les ambiguïtés de planning.")
bullet("H2 : L'affectation des activités avec priorités, échéances et progression renforce le suivi opérationnel.")
bullet("H3 : Messagerie et notifications diminuent les pertes d'information.")
bullet("H4 : Authentification sécurisée et isolation par entreprise assurent la confidentialité nécessaire.")
bullet("H5 : L'approche UP (Inception → Élaboration → Construction → Transition) améliore la maîtrise des risques et la qualité des livrables.")

h2("0.4 Objectifs de l'étude")
h3("0.4.1 Objectif général")
p("Concevoir et réaliser l'application web PLANIPRO en appliquant le Processus Unifié, afin de gérer de façon intégrée les rendez-vous et les activités professionnelles.")
h3("0.4.2 Objectifs spécifiques")
bullet("Élaborer un cadre théorique incluant le Processus Unifié, UML et les SI web.")
bullet("Réaliser l'étude de l'existant et la Vision du système (Inception).")
bullet("Spécifier les exigences et le modèle de cas d'utilisation (Élaboration).")
bullet("Concevoir l'architecture, le modèle de données et les interfaces.")
bullet("Implémenter les modules par itérations de Construction.")
bullet("Tester le système et documenter la Transition.")

h2("0.5 Intérêt du sujet")
p("Intérêt scientifique : contribution appliquée à l'usage du Processus Unifié pour un SI métier PME. Intérêt pratique : outil opérationnel déployable en local (XAMPP) ou adaptable à un hébergement. Intérêt pédagogique : synthèse d'analyse UML/UP, conception, développement PHP/MySQL et sécurité. Intérêt social : modernisation de la planification organisationnelle (ITU, 2023).")

h2("0.6 Délimitation du sujet")
p("Périmètre : application web (pas de native mobile) ; modules RDV, activités, contacts, messagerie, calendrier, notifications, rapports ; déploiement démo XAMPP ; rôles entreprise, visiteur (employe), client, partenaire et public. Hors périmètre version actuelle : facturation, visioconférence, SMTP avancé, API REST complète, SaaS multi-tenant commercial.")

h2("0.7 Méthodologie : le Processus Unifié (UP)")
p("Nous avons adopté le Processus Unifié (Unified Process) comme démarche de développement (Jacobson et al., 1999 ; Kruchten, 2003 ; Scott, 2001). UP se caractérise par quatre principes cardinaux :")
bullet("Piloté par les cas d'utilisation (use-case driven).")
bullet("Centré sur l'architecture (architecture-centric).")
bullet("Itératif et incrémental.")
bullet("Orienté risque (risk-driven).")
p("Les quatre phases structurantes sont : Inception (cadrage, Vision, risques majeurs), Élaboration (exigences, architecture de référence, cas d'utilisation prioritaires), Construction (construction itérative des fonctionnalités), Transition (recette, déploiement, formation). Ces phases sont détaillées au chapitre I et appliquées concrètement aux chapitres II à IV.")
figure("0.3", "Cycle de vie UP appliqué à PLANIPRO",
       "À dessiner : 4 phases (Inception, Élaboration, Construction, Transition) en colonnes + disciplines en lignes (style « bosse » RUP). Annoter PLANIPRO.")

h2("0.8 Subdivision du travail")
bullet("Chapitre I — Cadre théorique et conceptuel (SI, planification, web, BDD, sécurité, UP, UML).")
bullet("Chapitre II — Étude de l'existant et analyse du problème (Inception).")
bullet("Chapitre III — Analyse et conception (Élaboration).")
bullet("Chapitre IV — Implémentation (Construction et Transition).")
p("Une bibliographie abondante et des annexes complètent le document.", first_indent=1.25)
p("*** Fin de l'introduction générale ***", center=True, justify=False, italic=True, first_indent=0)

# =============================================================================
# CHAPITRE I
# =============================================================================
h1("Chapitre I — Cadre théorique et conceptuel")

h2("1.1 Introduction du chapitre")
p("Ce chapitre clarifie les notions mobilisées pour PLANIPRO : systèmes d'information, gestion des rendez-vous et activités, architectures web, bases de données, sécurité, et surtout le Processus Unifié et UML, qui constituent le pilier méthodologique du mémoire (Reix et al., 2016 ; Jacobson et al., 1999 ; Booch et al., 2005).")

h2("1.2 Les systèmes d'information")
h3("1.2.1 Définition")
p("Un système d'information (SI) est un ensemble organisé de ressources (humaines, matérielles, logicielles, données, procédures) permettant d'acquérir, traiter, stocker et diffuser de l'information (Reix, 2004 ; Davis & Olson, 1985). Alter (2008) insiste sur la notion de système de travail : le SI soutient les activités des acteurs.")
h3("1.2.2 Niveaux du SI")
p("On distingue classiquement les niveaux opérationnel, tactique et stratégique (Anthony, 1965 ; Laudon & Laudon, 2020). PLANIPRO se situe principalement aux niveaux opérationnel (transactions RDV/activités) et tactique (rapports, tableaux de bord).")
figure("1.1", "Pyramide des SI et positionnement de PLANIPRO",
       "Pyramide 3 niveaux + annotation PLANIPRO (base et milieu).")
table(
    ["Type", "Finalité", "Lien PLANIPRO"],
    [
        ["TPS", "Transactions quotidiennes", "Création / validation RDV"],
        ["MIS", "Rapports de gestion", "Module rapports"],
        ["CRM léger", "Relation client", "Contacts clients/partenaires"],
        ["Collaboration", "Coordination", "Activités + messagerie"],
    ],
    "Tableau 1.1 — Types de systèmes d'information (adapté de Laudon & Laudon, 2020)",
)
h3("1.2.3 Qualité de l'information")
p("Wang et Strong (1996) distinguent exactitude, complétude, cohérence, actualité et accessibilité. En planification, une erreur de créneau a un coût organisationnel élevé.")

h2("1.3 Gestion des rendez-vous et du temps organisationnel")
p("Le temps est une ressource rare (Drucker, 2006 ; Covey, 2004). Un rendez-vous professionnel est un engagement mutuel défini par un demandeur, un destinataire, un créneau, un statut et un motif (O'Brien & Marakas, 2011). Les systèmes de scheduling réduisent les erreurs et offrent une vision consolidée (Pinedo, 2016 ; Burke et al., 2004).")
figure("1.4", "Cycle de vie d'un rendez-vous dans PLANIPRO",
       "Automate : en_attente → accepte → termine ; refuse/annule ; reporte.")
p("PLANIPRO formalise les statuts : en_attente, accepte, refuse, reporte, annule, termine.")

h2("1.4 Gestion des activités et des tâches")
p("Une activité désigne ici une unité de travail interne assignable, priorisée (basse à urgente), munie d'un statut et d'une progression 0–100 % (PMI, 2021 ; Anderson, 2010). Cette approche s'inspire du management visuel Kanban sans imposer une méthodologie Agile complète.")

h2("1.5 Applications web et architectures logicielles")
p("Une application web est accessible via navigateur et communique en HTTP/HTTPS (Conallen, 2002). L'architecture à trois tiers sépare présentation, logique métier et données (Fowler, 2002).")
figure("1.2", "Architecture client-serveur web à trois tiers",
       "Navigateur ↔ Apache/PHP ↔ MySQL — annoter XAMPP.")
figure("1.3", "MVC et organisation pages/actions de PLANIPRO",
       "Schéma MVC + parallèle pages (vue) / actions (contrôle) / SQL (modèle).")
p("PLANIPRO adopte une organisation pragmatique pages/actions compatible avec le pattern Post-Redirect-Get. Le choix de PHP 8+ s'explique par sa disponibilité sur XAMPP et sa maturité pour les applications CRUD (Welling & Thomson, 2016 ; Nixon, 2021).")

h2("1.6 Bases de données relationnelles")
p("Le modèle relationnel (Codd, 1970) organise les données en tables liées par clés. La normalisation réduit redondances et anomalies (Date, 2003 ; Elmasri & Navathe, 2015). L'accès via PDO et requêtes préparées constitue une défense contre l'injection SQL (OWASP, 2021).")

h2("1.7 Sécurité des applications web")
p("Les risques majeurs incluent injection, XSS, CSRF, authentification faible et exposition de données (OWASP Top 10, 2021 ; Schneier, 2015). Contre-mesures retenues : password_hash/bcrypt, sessions, contrôle de rôle, jetons CSRF, échappement HTML, PDO, isolation multi-entreprises.")

h2("1.8 Le Processus Unifié (UP) : fondements et principes")
h3("1.8.1 Origine et définition")
p("Le Processus Unifié est un processus de développement logiciel itératif issu des travaux d'Ivar Jacobson, Grady Booch et James Rumbaugh, et industrialisé notamment sous la forme du Rational Unified Process (Jacobson et al., 1999 ; Kruchten, 2003 ; Kroll & Kruchten, 2003). Il fournit un cadre adaptable plutôt qu'une procédure rigide unique (Scott, 2001).")
h3("1.8.2 Caractéristiques fondamentales")
bullet("Use-case driven : les cas d'utilisation capturent les exigences fonctionnelles et guident l'analyse, la conception, l'implémentation et les tests (Cockburn, 2000).")
bullet("Architecture-centric : une architecture exécutable est construite tôt (Élaboration) pour stabiliser les choix structurants (Kruchten, 1995, 2003).")
bullet("Iterative and incremental : le système croît par itérations produisant des incréments testables (Larman, 2004).")
bullet("Risk-driven : les risques majeurs sont traités en priorité (Boehm, 1988).")
h3("1.8.3 Les quatre phases")
table(
    ["Phase", "Objectif principal", "Jalons typiques", "Application PLANIPRO"],
    [
        ["Inception", "Cadrer le projet, Vision, faisabilité", "Lifecycle Objective (LCO)", "Chap. II — existant, Vision, risques"],
        ["Élaboration", "Exigences, architecture de base", "Lifecycle Architecture (LCA)", "Chap. III — CU, architecture, BDD"],
        ["Construction", "Construire le produit par itérations", "Initial Operational Capability", "Chap. IV — modules PHP"],
        ["Transition", "Déployer, former, corriger", "Product Release", "Chap. IV — tests, installation, doc"],
    ],
    "Tableau 1.3 — Phases et objectifs du Processus Unifié",
)
figure("1.5", "Phases du Processus Unifié",
       "Frise Inception → Élaboration → Construction → Transition avec jalons LCO/LCA/IOC/PR.")
h3("1.8.4 Disciplines (workflows)")
p("UP organise le travail en disciplines : Modélisation métier, Exigences, Analyse & Conception, Implémentation, Test, Déploiement, Gestion de configuration, Gestion de projet, Environnement (Kruchten, 2003). L'intensité de chaque discipline varie selon la phase (courbes en « bosses »).")
figure("1.6", "Disciplines UP et artefacts pour PLANIPRO",
       "Tableau phases × disciplines ; cocher les artefacts produits (Vision, SRS/CU, modèle analyse, modèle conception, code, cas de test).")
table(
    ["Approche", "Nature", "Forces", "Limites"],
    [
        ["Cascade", "Séquentielle", "Simple à planifier", "Peu flexible"],
        ["Cycle en V", "Séquentiel + tests", "Traçabilité tests", "Rigidité"],
        ["Agile/Scrum", "Itératif léger", "Réactivité", "Moins d'accent architecture formelle"],
        ["UP / RUP", "Itératif + architecture + CU", "Risques, UML, jalons", "Peut être lourd si non adapté"],
    ],
    "Tableau 1.2 — Comparaison des approches de développement",
)
p("Pour un mémoire de conception d'application web métier, UP offre un excellent équilibre entre formalisme UML et itérations concrètes (Ambler, 2002 ; Larman, 2004).")

h2("1.9 UML et artefacts du Processus Unifié")
p("UML (Unified Modeling Language) est le langage de modélisation standard associé à UP (Booch et al., 2005 ; Rumbaugh et al., 2004 ; Roques, 2018). Artefacts mobilisés :")
table(
    ["Artefact UP", "Description", "Dans ce mémoire"],
    [
        ["Vision", "Objectifs, problèmes, parties prenantes", "Chap. II"],
        ["Glossaire", "Vocabulaire métier", "Annexe + §1.10"],
        ["Modèle de CU", "Acteurs et cas d'utilisation", "Chap. III"],
        ["Spécifications CU", "Scénarios nominaux/alternatifs", "Chap. III"],
        ["Modèle d'analyse", "Classes d'analyse, collaborations", "Chap. III"],
        ["Modèle de conception", "Architecture, BDD, UI", "Chap. III"],
        ["Modèle d'implémentation", "Composants, code", "Chap. IV"],
        ["Modèle de test", "Cas de test", "Chap. IV"],
        ["Document d'architecture logicielle", "Vues 4+1", "Chap. III"],
    ],
    "Tableau 1.4 — Artefacts UP mobilisés dans PLANIPRO",
)
figure("1.7", "Principe du pilotage par les cas d'utilisation",
       "Schéma : CU → Analyse → Conception → Implémentation → Tests (flèches de traçabilité).")
p("En francophonie, Merise (MCD/MLD/MPD) complète utilement UML pour la conception des données (Tardieu et al., 2000 ; Nanci & Espinasse, 2001). Nous combinons les deux.")

h2("1.10 Concepts clés de PLANIPRO")
bullet("Entreprise : organisation propriétaire d'un espace PLANIPRO.")
bullet("Visiteur (rôle technique employe) : compte opérationnel créé par l'entreprise.")
bullet("Client / Partenaire : interlocuteurs pouvant demander des RDV.")
bullet("Rendez-vous : engagement planifié à cycle de vie formalisé.")
bullet("Activité : tâche interne assignable et mesurable.")
bullet("Contact public : message depuis l'accueil sans compte obligatoire.")
bullet("Notification : alerte in-app ou publique (jeton).")
bullet("Conversation : messagerie 1–1 entre utilisateurs autorisés.")

h2("1.11 Conclusion du chapitre")
p("Le cadre théorique articule SI, planification, technologies web et, de façon centrale, le Processus Unifié et UML. Le chapitre suivant applique la phase d'Inception : étude de l'existant, Vision et analyse du problème.")
p("*** Fin du chapitre I ***", center=True, justify=False, italic=True, first_indent=0)

# =============================================================================
# CHAPITRE II — INCEPTION
# =============================================================================
h1("Chapitre II — Étude de l'existant et analyse du problème")

h2("2.1 Introduction (phase d'Inception)")
p("Dans le Processus Unifié, la phase d'Inception vise à comprendre le « quoi » et le « pourquoi » du projet : parties prenantes, problèmes, Vision, périmètre, risques majeurs et faisabilité (Kruchten, 2003 ; Kroll & Kruchten, 2003). Ce chapitre correspond à cette phase pour PLANIPRO.")

h2("2.2 Présentation du domaine d'étude")
p("Le domaine cible est celui des entreprises (PME, structures de services) qui planifient des interactions avec clients, partenaires et visiteurs. Elles doivent recevoir des demandes de rencontre, confirmer ou reporter des créneaux, répartir le travail interne, conserver une trace des échanges et disposer de synthèses.")
figure("2.1", "Cartographie des acteurs du domaine",
       "Entreprise au centre ; Clients, Partenaires, Visiteurs, Public ; flèches RDV/message/activité.")

h2("2.3 Modélisation métier de l'existant (as-is)")
p("La discipline UP de modélisation métier décrit les processus actuels (as-is) avant d'imaginer le to-be (Eriksson & Penker, 2000).")
h3("2.3.1 Processus actuel de prise de RDV")
bullet("Contact par téléphone, WhatsApp ou e-mail.", ordered=True, n=1)
bullet("Proposition orale d'un créneau après consultation d'un agenda personnel.", ordered=True, n=2)
bullet("Confirmation peu formalisée ; reports informels.", ordered=True, n=3)
bullet("Risques d'oubli, de conflit et de mauvaise transmission.", ordered=True, n=4)
bullet("Historique rarement consolidé.", ordered=True, n=5)
figure("2.2", "Processus as-is de prise de rendez-vous",
       "BPMN/activités : Client → Appel/WhatsApp → Secrétariat → Agenda → Confirmation orale. Points de rupture en rouge.")
table(
    ["Processus", "Support fréquent", "Formalisation"],
    [
        ["Demande de RDV", "Téléphone, WhatsApp, e-mail", "Faible"],
        ["Confirmation", "Oral / message", "Faible à moyenne"],
        ["Planning", "Agenda papier / Calendar perso", "Moyenne"],
        ["Activités", "Excel, post-it, mémoire", "Faible"],
        ["Historique / reporting", "Souvent inexistant", "Très faible"],
    ],
    "Tableau 2.1 — Inventaire des pratiques actuelles de planification",
)
h3("2.3.2 Activités internes et communication")
p("Les activités (relances, préparations, suivis) sont souvent non assignées formellement. Les échanges sur canaux personnels fragilisent la mémoire organisationnelle (ISO 9001, 2015 ; Allen, 2015).")

h2("2.4 Analyse critique et risques")
p("Faiblesses : fragmentation des outils, doubles réservations, absence de workflow de validation, faible traçabilité, mauvaise visibilité managériale, déperdition d'information. Forces : rapidité du contact informel, familiarité des outils, faible coût apparent (Davis, 1989 ; Venkatesh et al., 2003).")
figure("2.3", "Diagramme d'Ishikawa des dysfonctionnements",
       "Effet « Planification inefficace » ; causes Méthodes, Moyens, Main-d'œuvre, Milieu, Matériel, Management.")
figure("2.4", "Matrice SWOT de l'existant",
       "Forces / Faiblesses / Opportunités / Menaces.")
table(
    ["Risque", "Probabilité", "Impact", "Mitigation PLANIPRO"],
    [
        ["Double réservation", "Moyenne", "Élevé", "Planning central + statuts"],
        ["Oubli de RDV", "Élevée", "Élevé", "Calendrier + rappels 24 h"],
        ["Perte d'historique", "Élevée", "Moyen", "BDD + historique"],
        ["Fuite via téléphone perso", "Moyenne", "Élevé", "Messagerie + rôles"],
        ["Absence de reporting", "Élevée", "Moyen", "Rapports / PDF"],
    ],
    "Tableau 2.6 — Matrice des risques (Inception UP)",
)

h2("2.5 Formulation du problème et vision du système")
p("Problème central :", first_indent=0)
p("L'absence d'un SI intégré, multi-acteurs et sécurisé pour la gestion des rendez-vous et des activités entraîne fragmentation des données, erreurs de planning, faible traçabilité et communication difficilement auditable.",
  italic=True, first_indent=0, indent_left=1.0)
h3("2.5.1 Vision (artefact UP)")
p("Selon UP, le document de Vision décrit le problème, les parties prenantes, les fonctionnalités essentielles et les contraintes (Leffingwell & Widrig, 2003). Vision PLANIPRO : offrir une plateforme web simple où l'entreprise pilote RDV et activités, où clients/partenaires déposent des demandes, où les visiteurs suivent leurs tâches, et où le public peut entrer en contact avec suivi.")
figure("2.6", "Structure du document de Vision UP",
       "Blocs : Positionnement, Parties prenantes, Description problème, Fonctionnalités clés, Contraintes, Qualité.")

h2("2.6 Analyse comparative des solutions")
table(
    ["Critère", "Agenda/Excel", "Google Calendar", "Calendly", "Asana/Trello", "PLANIPRO"],
    [
        ["Workflow validation RDV", "Non/Partiel", "Partiel", "Partiel", "Faible", "Oui"],
        ["Activités + progression", "Partiel", "Non", "Non", "Oui", "Oui"],
        ["Messagerie métier", "Non", "Non", "Non", "Partiel", "Oui"],
        ["Contacts multi-rôles", "Partiel", "Partiel", "Partiel", "Partiel", "Oui"],
        ["Contact public + suivi", "Non", "Non", "Partiel", "Non", "Oui"],
        ["Hébergement local", "Oui", "Non", "Non", "Non", "Oui"],
    ],
    "Tableau 2.3 — Comparaison fonctionnelle des solutions",
)
figure("2.5", "Comparaison visuelle des outils",
       "Radar ou barres sur 6 axes ; mettre PLANIPRO en évidence.")

h2("2.7 Besoins et exigences initiales")
p("En Inception, UP produit une première version des exigences (Kruchten, 2003).")
table(
    ["ID", "Besoin fonctionnel", "Priorité"],
    [
        ["BF01", "Authentification multi-rôles", "Haute"],
        ["BF02", "Demander / valider / reporter des RDV", "Haute"],
        ["BF03", "Créer et assigner des activités", "Haute"],
        ["BF04", "Calendrier RDV + activités", "Haute"],
        ["BF05", "Gérer clients, partenaires, visiteurs", "Haute"],
        ["BF06", "Messagerie 1–1", "Moyenne"],
        ["BF07", "Contact public + notifications", "Moyenne"],
        ["BF08", "Rappels RDV", "Moyenne"],
        ["BF09", "Rapports / export PDF", "Moyenne"],
        ["BF10", "Paramètres profil / entreprise", "Basse"],
    ],
    "Tableau 2.4 — Besoins fonctionnels prioritaires",
)
table(
    ["ID", "Catégorie", "Exigence"],
    [
        ["BNF01", "Sécurité", "Hash, CSRF, rôles, isolation données"],
        ["BNF02", "Utilisabilité", "Navigation par rôle, formulaires clairs"],
        ["BNF03", "Performance", "Réponse acceptable en local"],
        ["BNF04", "Maintenabilité", "Code pages/actions/includes documenté"],
        ["BNF05", "Portabilité", "Stack PHP/MySQL/XAMPP"],
        ["BNF06", "Fiabilité", "Intégrité référentielle"],
    ],
    "Tableau 2.5 — Besoins non fonctionnels",
)
table(
    ["Forces existant", "Faiblesses existant"],
    [
        ["Rapidité du contact informel", "Pas de statuts formalisés"],
        ["Outils connus", "Données dispersées"],
        ["Coût direct faible", "Coûts cachés (oublis, stress)"],
        ["Flexibilité", "Reporting faible"],
    ],
    "Tableau 2.2 — Forces et faiblesses de l'existant",
)

h2("2.8 Conclusion du chapitre")
p("L'Inception a permis de diagnostiquer l'existant, d'identifier les risques, de formuler le problème et une Vision, et de dresser les exigences initiales. Le jalon Lifecycle Objective est atteint sur le plan académique : le projet est justifié et cadré. Le chapitre III développe la phase d'Élaboration.")
p("*** Fin du chapitre II ***", center=True, justify=False, italic=True, first_indent=0)

# =============================================================================
# CHAPITRE III — ÉLABORATION
# =============================================================================
h1("Chapitre III — Analyse et conception du système")

h2("3.1 Introduction (phase d'Élaboration)")
p("L'Élaboration UP a pour but de stabiliser les exigences majeures, d'élaborer une architecture de référence exécutable et de mitiger les risques principaux (Kruchten, 2003 ; Larman, 2004). Ce chapitre produit les artefacts d'exigences, d'analyse et de conception de PLANIPRO.")

h2("3.2 Exigences fonctionnelles et non fonctionnelles")
p("Les exigences précisent le cahier des charges issu du chapitre II (IEEE 830 ; Robertson & Robertson, 2012 ; Wiegers & Beatty, 2013). Le système doit permettre l'inscription entreprise, la connexion multi-rôles, le workflow complet des RDV, la gestion des activités, le calendrier, les contacts, la messagerie, le contact public, les notifications/rappels et les rapports.")
table(
    ["ID", "Exigence non fonctionnelle détaillée"],
    [
        ["ENF-S1", "Mots de passe hashés (password_hash)"],
        ["ENF-S2", "Jeton CSRF sur actions POST"],
        ["ENF-S3", "Contrôle require_login / require_role"],
        ["ENF-U1", "Interface responsive et menus par rôle"],
        ["ENF-P1", "Temps de réponse < 3 s écrans courants en local"],
        ["ENF-M1", "Séparation pages / actions / includes"],
    ],
    "Tableau 3.6 — Exigences non fonctionnelles détaillées",
)

h2("3.3 Acteurs et cas d'utilisation (UP)")
p("Le modèle de cas d'utilisation est l'artefact central d'UP (Jacobson et al., 1992 ; Cockburn, 2000).")
table(
    ["Acteur", "Rôle technique", "Création", "Responsabilités"],
    [
        ["Entreprise", "entreprise", "Inscription publique", "Piloter RDV, contacts, rapports"],
        ["Visiteur", "employe", "Créé par entreprise", "RDV/activités assignés, messagerie"],
        ["Client", "client", "Créé par entreprise", "Demander RDV, historique"],
        ["Partenaire", "partenaire", "Créé par entreprise", "Demander RDV"],
        ["Public", "—", "Sans compte", "Contacter, suivre notifs publiques"],
    ],
    "Tableau 3.1 — Acteurs du système PLANIPRO",
)
table(
    ["Fonction", "Entreprise", "Visiteur", "Client", "Partenaire", "Public"],
    [
        ["Valider RDV", "Oui", "—", "—", "—", "—"],
        ["Demander RDV", "Création", "—", "Oui", "Oui", "Via message"],
        ["Activités", "Oui", "Lecture", "—", "—", "—"],
        ["Calendrier", "Oui", "Ses événements", "—", "—", "—"],
        ["Messagerie", "Oui", "Oui", "Oui", "Oui", "—"],
        ["Rapports", "Oui", "—", "—", "—", "—"],
    ],
    "Tableau 3.2 — Matrice des droits d'accès",
)
figure("3.1", "Modèle de cas d'utilisation global PLANIPRO",
       "StarUML/Draw.io : acteurs + CU (Se connecter, Gérer RDV, Gérer activités, Contacts, Messagerie, Calendrier, Rapports, Contacter…).")
figure("3.2", "CU — module Rendez-vous",
       "Demander, Modifier, Annuler, Accepter, Refuser, Reporter, Terminer, Filtrer.")
figure("3.3", "CU — module Activités",
       "Créer, Modifier, Supprimer, Assigner, Progression, Consulter.")

h2("3.4 Scénarios et spécifications des CU")
h3("3.4.1 CU « Accepter un rendez-vous »")
p("Acteur principal : Entreprise. Préconditions : authentifiée ; RDV en_attente. Scénario nominal : ouvrir liste → choisir RDV → Accepter (+ commentaire) → statut accepte → historique + notification demandeur → événement calendrier. Exceptions : droits insuffisants ; RDV déjà traité.")
h3("3.4.2 CU « Demander un rendez-vous »")
p("Acteur : Client/Partenaire. Scénario : formulaire titre/motif/dates → insertion en_attente → notification entreprise.")
h3("3.4.3 CU « Contacter une entreprise » (Public)")
p("Scénario : formulaire public → contact_public + token → suivi notifications ; extension : création compte visiteur.")
figure("3.4", "Diagramme de séquence : demande et validation d'un RDV",
       "Client → UI → Action → BDD → Notif entreprise → rdv_action → Notif client → Calendrier.")
figure("3.5", "Contact public → compte visiteur",
       "Diagramme d'activités UML complet.")

h2("3.5 Analyse (modèle d'analyse)")
p("Le modèle d'analyse UP propose des classes d'analyse (boundary, control, entity) pour comprendre le domaine sans figer la technique (Jacobson et al., 1999 ; Larman, 2004).")
figure("3.6", "Modèle d'analyse (classes d'analyse)",
       "Entity : User, Entreprise, RendezVous, Activite, Message… ; Control : GestionRDV, Auth ; Boundary : écrans.")

h2("3.6 Conception de la base de données")
p("Douze tables principales : users, entreprises, clients, partenaires, rendez_vous, activites, notifications, historique, conversations, messages, contact_public, notifications_publiques.")
figure("3.7", "MCD / diagramme entité-association",
       "Looping/Draw.io avec cardinalités.")
figure("3.8", "MLD relationnel PLANIPRO",
       "Tables PK/FK et liens.")
table(
    ["Attribut", "Type", "Contraintes", "Description"],
    [
        ["id", "INT UNSIGNED", "PK, AI", "Identifiant"],
        ["entreprise_id", "INT UNSIGNED", "FK NOT NULL", "Entreprise"],
        ["demandeur_id", "INT UNSIGNED", "FK NOT NULL", "Demandeur"],
        ["type_demandeur", "ENUM", "NOT NULL", "client/partenaire/employe"],
        ["titre", "VARCHAR(200)", "NOT NULL", "Titre"],
        ["date_debut / date_fin", "DATETIME", "NOT NULL", "Créneau"],
        ["statut", "ENUM", "NOT NULL", "Cycle de vie"],
    ],
    "Tableau 3.3 — Dictionnaire de données (extrait rendez_vous)",
)
table(
    ["#", "Relation", "Cardinalité", "FK"],
    [
        ["1", "User possède Entreprise", "1 — 0..1", "entreprises.user_id"],
        ["2", "Entreprise rattache Users", "1 — 0..N", "users.entreprise_id"],
        ["3", "Entreprise reçoit RDV", "1 — 0..N", "rendez_vous.entreprise_id"],
        ["4", "User demande RDV", "1 — 0..N", "rendez_vous.demandeur_id"],
        ["5", "Conversation → Messages", "1 — 1..N", "messages.conversation_id"],
    ],
    "Tableau 3.4 — Relations et cardinalités (extrait)",
)

h2("3.7 Architecture logicielle (vues 4+1)")
p("Kruchten (1995) propose les vues 4+1 : cas d'utilisation (+1), logique, processus, développement, déploiement. PLANIPRO : monolithique web PHP, pattern pages/actions, MySQL InnoDB, assets publics, uploads logos.")
figure("3.9", "Architecture 4+1 appliquée à PLANIPRO",
       "5 vues annotées pour le projet.")
p("Principes : PRG ; PDO préparé ; échappement HTML ; isolation par entreprise_id / demandeur_id / assigne_a.")

h2("3.8 Conception des interfaces")
p("Conception guidée par les heuristiques de Nielsen (1994) et une navigation par rôle.")
figure("3.10", "Maquette page d'accueil publique", "Wireframe hero + entreprises + CTA.")
figure("3.11", "Maquette tableau de bord entreprise", "KPI + raccourcis.")
figure("3.12", "Maquette calendrier", "FullCalendar mois/semaine/jour.")
figure("3.13", "Maquette messagerie", "Liste + fil + saisie.")

h2("3.9 Sécurité et règles de gestion")
table(
    ["ID", "Règle de gestion"],
    [
        ["RG01", "Seules les entreprises s'inscrivent publiquement"],
        ["RG02", "Clients/partenaires/visiteurs créés par l'entreprise"],
        ["RG03", "Un client ne demande des RDV qu'à son entreprise"],
        ["RG04", "Seule l'entreprise valide les statuts des RDV reçus"],
        ["RG05", "Un visiteur ne voit que ses RDV/activités"],
        ["RG06", "Conversation 1–1 unique par paire d'utilisateurs"],
        ["RG07", "Suivi public par token sans exposer les autres données"],
        ["RG08", "Mots de passe jamais en clair"],
        ["RG09", "POST sensibles protégés par CSRF"],
        ["RG10", "Comptes désactivables (actif=0)"],
    ],
    "Tableau 3.5 — Règles de gestion métier",
)

h2("3.10 Conclusion du chapitre")
p("L'Élaboration a produit exigences, modèle de CU, analyse, conception données/architecture/UI et règles de sécurité. Le jalon Lifecycle Architecture est atteint : l'architecture de référence est définie. La Construction fait l'objet du chapitre IV.")
p("*** Fin du chapitre III ***", center=True, justify=False, italic=True, first_indent=0)

# =============================================================================
# CHAPITRE IV — CONSTRUCTION + TRANSITION
# =============================================================================
h1("Chapitre IV — Implémentation du système proposé")

h2("4.1 Introduction (Construction et Transition)")
p("La Construction UP construit le système par itérations produisant des incréments exécutables ; la Transition prépare le déploiement, la recette et le support (Kruchten, 2003 ; McConnell, 2004).")

h2("4.2 Environnement et organisation du code")
table(
    ["Composant", "Outil", "Rôle"],
    [
        ["OS", "Windows 10", "Développement"],
        ["Serveur", "Apache (XAMPP)", "HTTP"],
        ["SGBD", "MySQL/MariaDB", "Persistance"],
        ["Langage", "PHP 8+", "Métier"],
        ["Front", "HTML/CSS/JS + FullCalendar", "UI"],
        ["Admin BDD", "phpMyAdmin", "Import SQL"],
    ],
    "Tableau 4.1 — Environnement logiciel",
)
table(
    ["Répertoire", "Responsabilité"],
    [
        ["index.php", "Accueil public"],
        ["config/", "Connexion BDD"],
        ["includes/", "Auth, helpers, messagerie, rappels"],
        ["pages/", "Interfaces par rôle"],
        ["actions/", "Traitements POST"],
        ["public/", "CSS, JS"],
        ["sql/", "Schéma et seeds"],
        ["uploads/", "Logos"],
    ],
    "Tableau 4.2 — Structure des répertoires",
)
figure("4.11", "Schéma de déploiement local (XAMPP)",
       "Machine → XAMPP → htdocs/PLANIPRO → localhost.")

h2("4.3 Itérations de construction")
p("Conformément à UP, la Construction a été découpée en itérations à valeur croissante (Larman, 2004) :")
table(
    ["Itération", "Objectif", "Incrément livré", "Risques traités"],
    [
        ["C1", "Socle", "BDD + auth + layouts", "Connexion, rôles"],
        ["C2", "RDV", "Demande + validation + notifs", "Workflow métier"],
        ["C3", "Activités + calendrier", "CRUD activités + FullCalendar", "Agrégation événements"],
        ["C4", "Contacts & visiteurs", "Clients/partenaires/employés", "Rattachement"],
        ["C5", "Messagerie & public", "Conversations + contact_public", "Token public"],
        ["C6", "Rapports & polish", "PDF, UX, correctifs", "Transition"],
    ],
    "Tableau 4.3 — Itérations de Construction UP",
)
figure("4.12", "Plan d'itérations UP (Construction)",
       "Timeline C1…C6 avec livrables.")

h2("4.4 Implémentation des modules")
h3("4.4.1 Authentification")
p("Sessions PHP, password_verify, role_home(), require_login/require_role, CSRF, inscription entreprise créant users + entreprises.")
figure("4.2", "Inscription et connexion", "Captures d'écran réelles.")
h3("4.4.2 Rendez-vous")
p("Actions rdv_demande, rdv_action, rdv_entreprise_save ; transitions de statut ; historique et notifications.")
figure("4.4", "Gestion des rendez-vous", "Liste + badges de statut + actions.")
h3("4.4.3 Activités")
p("activite_save, priorités, progression, assignation, échéances vers calendrier.")
figure("4.5", "Formulaire et liste d'activités", "Capture.")
h3("4.4.4 Calendrier")
p("Endpoint JSON calendrier_events.php consommé par FullCalendar (vues mois/semaine/jour/liste).")
figure("4.6", "Calendrier vue mois", "Capture avec légende couleurs.")
h3("4.4.5 Contacts, messagerie, public, rapports")
p("Gestion contacts/visiteurs ; conversations 1–1 ; contact public + notifications_publiques ; export PDF (SimplePdf).")
figure("4.7", "Messagerie intégrée", "Capture.")
figure("4.8", "Rapports / export PDF", "Capture.")
figure("4.10", "Notifications publiques", "Capture.")

h2("4.5 Interfaces réalisées")
figure("4.1", "Page d'accueil PLANIPRO", "Capture hero + entreprises.")
figure("4.3", "Tableau de bord entreprise", "Capture KPI.")
figure("4.9", "Espace client", "Demande RDV / historique.")

h2("4.6 Tests et transition")
p("La discipline Test d'UP assure la vérification des incréments (ISTQB, 2019 ; Kaner et al., 1999). La Transition couvre installation, comptes démo, documentation README et formation minimale.")
table(
    ["ID", "Scénario", "Résultat", "Statut"],
    [
        ["T01", "Inscription entreprise", "Compte créé", "OK"],
        ["T02", "Mauvais mot de passe", "Erreur, pas de session", "OK"],
        ["T03", "Demande RDV client", "en_attente + notif", "OK"],
        ["T04", "Acceptation RDV", "accepte + calendrier", "OK"],
        ["T05", "Activité assignée", "Visible visiteur", "OK"],
        ["T06", "Isolation multi-entreprises", "Accès refusé", "OK"],
        ["T07", "POST sans CSRF", "Rejet", "OK"],
        ["T08", "Contact public", "Suivi par token", "OK"],
        ["T09", "Export PDF", "Généré", "OK"],
        ["T10", "Rappel < 24 h", "Notification", "OK"],
    ],
    "Tableau 4.4 — Scénarios de test",
)
table(
    ["Module", "Conçu", "Implémenté", "Testé"],
    [
        ["Authentification", "Oui", "Oui", "Oui"],
        ["Rendez-vous", "Oui", "Oui", "Oui"],
        ["Activités", "Oui", "Oui", "Oui"],
        ["Calendrier", "Oui", "Oui", "Oui"],
        ["Contacts", "Oui", "Oui", "Oui"],
        ["Messagerie", "Oui", "Oui", "Oui"],
        ["Notifications", "Oui", "Oui", "Oui"],
        ["Rapports", "Oui", "Oui", "Oui"],
        ["Contact public", "Oui", "Oui", "Oui"],
    ],
    "Tableau 4.5 — Couverture fonctionnelle",
)
table(
    ["Critère", "As-is", "To-be (PLANIPRO)"],
    [
        ["Planning", "Outils personnels", "Plateforme centralisée"],
        ["Statuts RDV", "Implicites", "Workflow 6 états"],
        ["Activités", "Post-it / Excel", "Assignation + progression"],
        ["Communication", "Canaux perso", "Messagerie + notifs"],
        ["Pilotage", "Faible", "Dashboard + PDF"],
        ["Sécurité", "Faible", "Rôles, CSRF, hash"],
    ],
    "Tableau 4.6 — Passage as-is → to-be",
)
p("Les hypothèses H1–H5 trouvent un appui qualitatif dans les tests et le déroulé UP.")

h2("4.7 Difficultés, limites et perspectives")
p("Difficultés : isolation multi-entreprises ; transitions RDV ; token public ; agrégation calendrier ; encodage UTF-8. Limites : pas de SMTP reset password, pas de pièces jointes, pas de créneaux bloqués avancés, pas d'API REST, déploiement démo local. Perspectives : e-mail transactionnel, API mobile, disponibilités, migrations versionnées, HTTPS production, analytique enrichie.")
p("Recommandations Transition/production : HTTPS, secrets forts, sauvegardes, UAT 2–4 semaines, formation utilisateurs (OWASP, 2021 ; Bass et al., 2012).")

h2("4.8 Conclusion du chapitre")
p("Construction et Transition ont produit un prototype cohérent avec l'Élaboration, itéré selon UP, testé et documenté. Le jalon Product Release académique est atteint sous forme de prototype démontrable.")
p("*** Fin du chapitre IV ***", center=True, justify=False, italic=True, first_indent=0)

# =============================================================================
# CONCLUSION
# =============================================================================
h1("Conclusion générale et perspectives")
p("La problématique était de concevoir et réaliser, selon le Processus Unifié, une application web centralisant rendez-vous, activités et échanges professionnels. Le chapitre I a posé le socle théorique (SI, web, UP, UML). Le chapitre II (Inception) a diagnostiqué l'existant et formulé la Vision. Le chapitre III (Élaboration) a spécifié CU, architecture et données. Le chapitre IV (Construction/Transition) a livré et testé PLANIPRO.")
p("Les objectifs sont atteints et les hypothèses trouvent un appui qualitatif. UP s'est révélé adapté : pilotage par les cas d'utilisation, architecture clarifiée tôt, itérations maîtrisables, risques traités progressivement (Jacobson et al., 1999 ; Kruchten, 2003).")
p("Perspectives : industrialisation (SMTP, API, HTTPS), évaluation empirique en entreprise (Hevner et al., 2004), enrichissement analytique. PLANIPRO constitue un noyau réutilisable pour la digitalisation de la planification en contexte PME.")

# =============================================================================
# BIBLIOGRAPHIE (abondante, focus UP)
# =============================================================================
h1("Bibliographie et webographie")
p("Les références suivent un style inspiré de APA (7e éd.).", italic=True, first_indent=0, size=11)

h2("A. Processus Unifié, RUP, UML et génie logiciel")
refs_up = [
    "Ambler, S. W. (2002). Agile Modeling. Wiley.",
    "Ambler, S. W. (2004). The Object Primer (3e éd.). Cambridge University Press.",
    "Beck, K., et al. (2001). Manifesto for Agile Software Development. https://agilemanifesto.org",
    "Boehm, B. W. (1981). Software Engineering Economics. Prentice Hall.",
    "Boehm, B. W. (1988). A spiral model of software development and enhancement. IEEE Computer, 21(5), 61–72.",
    "Booch, G., Rumbaugh, J., & Jacobson, I. (2005). The Unified Modeling Language User Guide (2e éd.). Addison-Wesley.",
    "Brooks, F. P. (1995). The Mythical Man-Month (éd. anniversaire). Addison-Wesley.",
    "Cockburn, A. (2000). Writing Effective Use Cases. Addison-Wesley.",
    "Cohn, M. (2004). User Stories Applied. Addison-Wesley.",
    "Conallen, J. (2002). Building Web Applications with UML (2e éd.). Addison-Wesley.",
    "Dennis, A., Wixom, B. H., & Tegarden, D. (2015). Systems Analysis and Design with UML (5e éd.). Wiley.",
    "Fowler, M. (2002). Patterns of Enterprise Application Architecture. Addison-Wesley.",
    "Fowler, M. (2018). Refactoring (2e éd.). Addison-Wesley.",
    "Gamma, E., Helm, R., Johnson, R., & Vlissides, J. (1995). Design Patterns. Addison-Wesley.",
    "IEEE Computer Society. (2014). SWEBOK Guide (v3).",
    "IEEE. (1998). IEEE Std 830 — Software Requirements Specifications.",
    "Jacobson, I., Booch, G., & Rumbaugh, J. (1999). The Unified Software Development Process. Addison-Wesley.",
    "Jacobson, I., Christerson, M., Jonsson, P., & Övergaard, G. (1992). Object-Oriented Software Engineering: A Use Case Driven Approach. Addison-Wesley.",
    "Kroll, P., & Kruchten, P. (2003). The Rational Unified Process Made Easy. Addison-Wesley.",
    "Kruchten, P. (1995). The 4+1 view model of architecture. IEEE Software, 12(6), 42–50.",
    "Kruchten, P. (2003). The Rational Unified Process: An Introduction (3e éd.). Addison-Wesley.",
    "Larman, C. (2004). Agile and Iterative Development: A Manager's Guide. Addison-Wesley.",
    "Larman, C. (2005). Applying UML and Patterns (3e éd.). Prentice Hall.",
    "Leffingwell, D., & Widrig, D. (2003). Managing Software Requirements (2e éd.). Addison-Wesley.",
    "Martin, R. C. (2008). Clean Code. Prentice Hall.",
    "Martin, R. C. (2017). Clean Architecture. Prentice Hall.",
    "McConnell, S. (2004). Code Complete (2e éd.). Microsoft Press.",
    "Övergaard, G., & Palmkvist, K. (2004). Use Cases: Patterns and Blueprints. Addison-Wesley.",
    "Pohl, K. (2010). Requirements Engineering. Springer.",
    "Pressman, R. S., & Maxim, B. R. (2019). Software Engineering (9e éd.). McGraw-Hill.",
    "Robertson, S., & Robertson, J. (2012). Mastering the Requirements Process (3e éd.). Addison-Wesley.",
    "Roques, P. (2018). UML 2.5 par la pratique (8e éd.). Eyrolles.",
    "Rumbaugh, J., Jacobson, I., & Booch, G. (2004). The Unified Modeling Language Reference Manual (2e éd.). Addison-Wesley.",
    "Schwaber, K., & Sutherland, J. (2020). The Scrum Guide. https://scrumguides.org",
    "Scott, K. (2001). The Unified Process Explained. Addison-Wesley.",
    "Sommerville, I. (2015). Software Engineering (10e éd.). Pearson.",
    "Wiegers, K., & Beatty, J. (2013). Software Requirements (3e éd.). Microsoft Press.",
]
for r in refs_up:
    ref(r)

h2("B. Systèmes d'information et management")
refs_si = [
    "Alter, S. (2008). Defining information systems as work systems. European Journal of Information Systems, 17(5), 448–469.",
    "Anthony, R. N. (1965). Planning and Control Systems. Harvard Business School.",
    "Avison, D., & Fitzgerald, G. (2006). Information Systems Development (4e éd.). McGraw-Hill.",
    "Checkland, P., & Holwell, S. (1998). Information, Systems and Information Systems. Wiley.",
    "Davenport, T. H. (1993). Process Innovation. Harvard Business School Press.",
    "Davis, F. D. (1989). Perceived usefulness, perceived ease of use… MIS Quarterly, 13(3), 319–340.",
    "Davis, G. B., & Olson, M. H. (1985). Management Information Systems (2e éd.). McGraw-Hill.",
    "DeLone, W. H., & McLean, E. R. (2003). The DeLone and McLean model… JMIS, 19(4), 9–30.",
    "Drucker, P. F. (2006). The Effective Executive. HarperBusiness.",
    "Hammer, M., & Champy, J. (2003). Reengineering the Corporation. HarperBusiness.",
    "Henderson, J. C., & Venkatraman, N. (1993). Strategic alignment… IBM Systems Journal, 32(1), 4–16.",
    "Hevner, A. R., et al. (2004). Design science in IS research. MIS Quarterly, 28(1), 75–105.",
    "Kaplan, R. S., & Norton, D. P. (1996). The Balanced Scorecard. Harvard Business School Press.",
    "Laudon, K. C., & Laudon, J. P. (2020). Management Information Systems (16e éd.). Pearson.",
    "Markus, M. L. (1983). Power, politics, and MIS implementation. CACM, 26(6), 430–444.",
    "Mintzberg, H. (1994). The Rise and Fall of Strategic Planning. Free Press.",
    "Nonaka, I., & Takeuchi, H. (1995). The Knowledge-Creating Company. Oxford University Press.",
    "O'Brien, J. A., & Marakas, G. M. (2011). Management Information Systems (10e éd.). McGraw-Hill.",
    "Pearlson, K. E., Saunders, C. S., & Galletta, D. F. (2019). Managing and Using Information Systems (7e éd.). Wiley.",
    "Porter, M. E. (1985). Competitive Advantage. Free Press.",
    "Reix, R. (2004). Systèmes d'information et management des organisations (5e éd.). Vuibert.",
    "Reix, R., Fallery, B., Kalika, M., & Rowe, F. (2016). Systèmes d'information et management (7e éd.). Vuibert.",
    "Seddon, P. B. (1997). A respecification… of the DeLone and McLean model. ISR, 8(3), 240–253.",
    "Stair, R., & Reynolds, G. (2017). Principles of Information Systems (13e éd.). Cengage.",
    "Turban, E., Pollard, C., & Wood, G. (2018). Information Technology for Management (11e éd.). Wiley.",
    "Venkatesh, V., et al. (2003). User acceptance of information technology… MIS Quarterly, 27(3), 425–478.",
    "Wang, R. Y., & Strong, D. M. (1996). Beyond accuracy… JMIS, 12(4), 5–33.",
    "Weill, P., & Ross, J. W. (2004). IT Governance. Harvard Business School Press.",
    "Zachman, J. A. (1987). A framework for information systems architecture. IBM Systems Journal, 26(3), 276–292.",
]
for r in refs_si:
    ref(r)

h2("C. Bases de données, web, sécurité, planification")
refs_tech = [
    "Anderson, D. J. (2010). Kanban. Blue Hole Press.",
    "Anderson, R. (2020). Security Engineering (3e éd.). Wiley.",
    "Bass, L., Clements, P., & Kazman, R. (2012). Software Architecture in Practice (3e éd.). Addison-Wesley.",
    "Burke, E. K., et al. (2004). The state of the art of nurse rostering. Journal of Scheduling, 7(6), 441–499.",
    "Chamberlin, D. D., & Boyce, R. F. (1974). SEQUEL… ACM SIGFIDET, 249–264.",
    "Chen, P. P. (1976). The entity-relationship model. ACM TODS, 1(1), 9–36.",
    "Codd, E. F. (1970). A relational model of data… CACM, 13(6), 377–387.",
    "Connolly, T., & Begg, C. (2014). Database Systems (6e éd.). Pearson.",
    "Covey, S. R. (2004). The 7 Habits of Highly Effective People. Free Press.",
    "Date, C. J. (2003). An Introduction to Database Systems (8e éd.). Addison-Wesley.",
    "Elmasri, R., & Navathe, S. B. (2015). Fundamentals of Database Systems (7e éd.). Pearson.",
    "Eriksson, H.-E., & Penker, M. (2000). Business Modeling with UML. Wiley.",
    "Fielding, R. T. (2000). Architectural Styles… (PhD). UC Irvine.",
    "Gabay, J., & Gabay, D. (2008). Merise et UML. Dunod.",
    "ISO. (2015). ISO 9001:2015 Quality management systems.",
    "ISO/IEC. (2011). ISO/IEC 25010 SQuaRE.",
    "ISTQB. (2019). Certified Tester Foundation Level Syllabus.",
    "Kaner, C., Falk, J., & Nguyen, H. Q. (1999). Testing Computer Software (2e éd.). Wiley.",
    "Kerzner, H. (2017). Project Management (12e éd.). Wiley.",
    "Nanci, D., & Espinasse, B. (2001). Ingénierie des SI : Merise. Vuibert.",
    "Nielsen, J. (1994). Usability Engineering. Morgan Kaufmann.",
    "Nixon, R. (2021). Learning PHP, MySQL & JavaScript (6e éd.). O'Reilly.",
    "Norman, D. A. (2013). The Design of Everyday Things. Basic Books.",
    "OWASP Foundation. (2021). OWASP Top 10. https://owasp.org/Top10/",
    "Pinedo, M. L. (2016). Scheduling (5e éd.). Springer.",
    "PMI. (2021). PMBOK Guide (7e éd.).",
    "Provos, N., & Mazières, D. (1999). A future-adaptable password scheme. USENIX ATC.",
    "Schneier, B. (2015). Data and Goliath. W. W. Norton.",
    "Shneiderman, B., et al. (2016). Designing the User Interface (6e éd.). Pearson.",
    "Silberschatz, A., Korth, H. F., & Sudarshan, S. (2019). Database System Concepts (7e éd.). McGraw-Hill.",
    "Tardieu, H., Rochfeld, A., & Colletti, R. (2000). La méthode Merise. Éditions d'Organisation.",
    "Welling, L., & Thomson, L. (2016). PHP and MySQL Web Development (5e éd.). Addison-Wesley.",
    "Allen, D. (2015). Getting Things Done. Penguin.",
    "Brynjolfsson, E., & McAfee, A. (2014). The Second Machine Age. Norton.",
    "Westerman, G., Bonnet, D., & McAfee, A. (2014). Leading Digital. HBR Press.",
]
for r in refs_tech:
    ref(r)

h2("D. Rapports, normes et webographie technique")
refs_web = [
    "Banque mondiale. (2021). World Development Report 2021: Data for Better Lives.",
    "Gartner. (2022). Market Guide for Meeting / Scheduling Solutions.",
    "GSMA. (2023). The Mobile Economy.",
    "ITU. (2023). Measuring digital development: Facts and figures.",
    "McKinsey Global Institute. (2016). Digital globalization.",
    "NIST. (2017). Digital Identity Guidelines (SP 800-63).",
    "Apache Friends. (2024). XAMPP Documentation. https://www.apachefriends.org",
    "FullCalendar. (2024). Documentation. https://fullcalendar.io/docs",
    "MySQL. (2024). MySQL 8.0 Reference Manual. https://dev.mysql.com/doc/",
    "PHP Group. (2024). PHP Manual. https://www.php.net/manual/",
    "MDN Web Docs. (2024). https://developer.mozilla.org",
    "MariaDB Foundation. (2024). https://mariadb.com/kb/",
    "Calendly Help. (2024). https://help.calendly.com",
    "Google Calendar Help. (2024). https://support.google.com/calendar",
    "diagrams.net. (2024). https://www.diagrams.net/doc",
    "StarUML Docs. (2024). https://docs.staruml.io",
    "W3C. (2023). WCAG 2.2.",
    "IETF. (2022). RFC 9110 — HTTP Semantics.",
    "PLANIPRO. (2024–2025). README.md ; docs/CONCEPTION.md ; sql/schema.sql.",
]
for r in refs_web:
    ref(r)

h2("E. Méthodologie de recherche et compléments")
refs_meth = [
    "Benbasat, I., Goldstein, D. K., & Mead, M. (1987). The case research strategy… MIS Quarterly, 11(3), 369–386.",
    "Cadle, J., Paul, D., & Turner, P. (2014). Business Analysis Techniques (2e éd.). BCS.",
    "Miles, M. B., Huberman, A. M., & Saldaña, J. (2014). Qualitative Data Analysis (3e éd.). Sage.",
    "Yin, R. K. (2018). Case Study Research and Applications (6e éd.). Sage.",
    "Zikmund, W. G., et al. (2013). Business Research Methods (9e éd.). Cengage.",
    "Baskerville, R., & Myers, M. D. (2002). Information systems as a reference discipline. MIS Quarterly, 26(1), 1–14.",
    "Lyytinen, K., & Hirschheim, R. (1987). Information systems failures… Oxford Surveys in IT, 4, 257–309.",
    "Orlikowski, W. J. (1992). The duality of technology. Organization Science, 3(3), 398–427.",
    "Parasuraman, A., Zeithaml, V. A., & Berry, L. L. (1988). SERVQUAL. Journal of Retailing, 64(1), 12–40.",
    "Rogers, E. M. (2003). Diffusion of Innovations (5e éd.). Free Press.",
    "Hoffer, J. A., George, J., & Valacich, J. (2016). Modern Systems Analysis and Design (8e éd.). Pearson.",
    "Kendall, K. E., & Kendall, J. E. (2013). Systems Analysis and Design (9e éd.). Pearson.",
    "Whitten, J. L., & Bentley, L. D. (2007). Systems Analysis and Design Methods (7e éd.). McGraw-Hill.",
    "Morley, C., et al. (2011). Processus métier et systèmes d'information. Dunod.",
    "Preece, J., Rogers, Y., & Sharp, H. (2015). Interaction Design (4e éd.). Wiley.",
]
for r in refs_meth:
    ref(r)

# =============================================================================
# ANNEXES
# =============================================================================
h1("Annexes")
h2("Annexe A — Installation rapide")
bullet("Démarrer Apache + MySQL (XAMPP).", ordered=True, n=1)
bullet("Importer sql/schema.sql puis sql/seed.sql.", ordered=True, n=2)
bullet("Ouvrir http://localhost/PLANIPRO/.", ordered=True, n=3)
p("Comptes démo (mot de passe : password) : entreprise@planipro.local, employe@planipro.local, client@planipro.local, partenaire@planipro.local.", first_indent=0)

h2("Annexe B — Extrait SQL")
figure("B.1", "Extrait schema.sql", "Coller CREATE TABLE users, entreprises, rendez_vous, activites.")

h2("Annexe C — Dictionnaire de données complet")
figure("C.1", "12 tables détaillées", "Reprendre docs/CONCEPTION.md.")

h2("Annexe D — Captures d'écran complémentaires")
figure("D.1", "Captures additionnelles", "Profil, paramètres, création visiteur, etc.")

h2("Annexe E — Matrice de traçabilité UP (extrait)")
table(
    ["Exigence / CU", "Analyse", "Conception", "Implémentation", "Test"],
    [
        ["BF02 / CU RDV", "§3.4", "Tables RDV", "rdv_*.php", "T03–T04"],
        ["BF03 / CU Activités", "§3.3", "Table activites", "activite_*.php", "T05"],
        ["BNF01 Sécurité", "§3.9", "Auth design", "auth.php", "T02,T06,T07"],
    ],
    "Tableau Annexe E — Traçabilité exigences → tests (esprit UP)",
)

h2("Annexe F — Glossaire")
bullet("UP : Processus Unifié de développement logiciel.")
bullet("Cas d'utilisation : unité d'exigence fonctionnelle vue acteur.")
bullet("Inception / Élaboration / Construction / Transition : phases UP.")
bullet("Visiteur : compte opérationnel (rôle employe) créé par l'entreprise.")

h2("Annexe G — Checklist soutenance")
bullet("Personnaliser page de garde (noms, logos).")
bullet("Remplacer tous les [EMPLACEMENT FIGURE] par images réelles.")
bullet("Mettre à jour la table des matières (références croisées Word).")
bullet("Vérifier la bibliographie citée dans le texte.")

p("")
p("— FIN DU MÉMOIRE —", center=True, justify=False, bold=True, size=14, first_indent=0)
p("Méthodologie : Processus Unifié (UP) — Projet PLANIPRO", center=True, justify=False, italic=True, size=11, first_indent=0)
p("Fichier : documents/MEMOIRE_PLANIPRO.docx", center=True, justify=False, size=10, first_indent=0)

doc.save(OUT)
print("OK:", OUT)
print("Size_KB:", round(OUT.stat().st_size / 1024, 1))
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
